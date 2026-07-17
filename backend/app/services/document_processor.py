import logging

logger = logging.getLogger(__name__)


def process_document(file_path: str, ext: str) -> str:
    try:
        if ext == ".txt":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()

        elif ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            return "\n".join([page.extract_text() or "" for page in reader.pages])

        elif ext == ".docx":
            from docx import Document as DocxDocument
            doc = DocxDocument(file_path)
            return "\n".join([p.text for p in doc.paragraphs])

        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(file_path)
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        texts.append(shape.text)
            return "\n".join(texts)

        elif ext in (".png", ".jpg", ".jpeg"):
            try:
                from PIL import Image
                import pytesseract
                image = Image.open(file_path)
                return pytesseract.image_to_string(image)
            except ImportError:
                return "[OCR not available. Install pytesseract and tesseract-ocr]"
            except Exception as e:
                logger.error(f"OCR failed: {e}")
                return f"[OCR processing error: {e}]"

        return "[Unsupported file type]"
    except Exception as e:
        logger.error(f"Document processing error for {file_path}: {e}")
        return f"[Error processing document: {e}]"
